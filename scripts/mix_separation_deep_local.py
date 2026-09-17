#!/usr/bin/env python3
"""
Глубокие тест-кейсы разбора миксов, по каждому этапу конвейера:

  1. parse_sklad_template        — разбор заполненного шаблона (спрос по длинам)
  2. scan_mix_pool / стебли      — чтение остатка Mix из Biflorica (в т.ч. сплит-строк)
  3. plan_mix_allocation         — распределение спроса по строкам Mix (жадно по цене)
  4. apply_mix_plans_to_biflorica — запись изменений в файл (списание/удаление строк)
  5. run_mix_separation          — путь целиком, от файлов до результата

Каждый тест создаёт синтетические .xlsx во временной папке — реальные файлы
из testdata/ не трогает. Запуск из корня репо:

  .venv/bin/python scripts/mix_separation_deep_local.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import traceback
from collections.abc import Callable
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook  # noqa: E402

from cvetopt.invoice.biflorica_mixes import (  # noqa: E402
    apply_mix_plans_to_biflorica,
    parse_sklad_template,
    plan_mix_allocation,
    run_mix_separation,
    scan_mix_pool,
    _as_qty,
    _biflorica_header,
    _priced_length_labels,
    _stems_for_length,
)
from cvetopt.invoice.xlsx_read import grid_by_row, read_excel_grid  # noqa: E402

# --- Структура тестовых файлов (сверена с реальными: testdata/mixes/*.xlsx,
#     "шаблоны тест/*.xlsx") ---
BIF_LENGTH_COL = {"40": "E", "50": "F", "60": "G", "70": "H", "80": "I",
                  "90": "J", "100": "K", "100+": "L"}
BIF_HEADER_ROW = 6
TPL_LENGTH_COL = {"50": "F", "60": "G", "70": "H", "80": "I"}
TPL_HEADER_ROW = 8

TESTS: list[tuple[str, Callable[[Path], None]]] = []


def test(name: str) -> Callable[[Callable[[Path], None]], Callable[[Path], None]]:
    def deco(fn: Callable[[Path], None]) -> Callable[[Path], None]:
        TESTS.append((name, fn))
        return fn

    return deco


# --- Хелперы для генерации синтетических файлов ---


def make_biflorica(path: Path, rows: list[dict]) -> Path:
    """rows: {plantation, type, variety, prices {length: price}, stems (int|str)}."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Отчёт по сделкам"
    h = BIF_HEADER_ROW
    ws[f"A{h}"] = "ДАТА И ВРЕМЯ СДЕЛКИ"
    ws[f"B{h}"] = "ПЛАНТАЦИЯ"
    ws[f"C{h}"] = "ТИП"
    ws[f"D{h}"] = "СОРТ"
    for lab, col in BIF_LENGTH_COL.items():
        ws[f"{col}{h}"] = lab
    ws[f"M{h}"] = "КОЛ-ВО ЯЩ."
    ws[f"N{h}"] = "ТИП ЯЩ."
    ws[f"O{h}"] = "ВСЕГО СТЕБЛЕЙ"
    ws[f"P{h}"] = "СУММА СДЕЛКИ"

    for i, r in enumerate(rows):
        rn = h + 1 + i
        ws[f"A{rn}"] = r.get("date", "28 May 2026, 10:34")
        ws[f"B{rn}"] = r.get("plantation", "PLANT")
        ws[f"C{rn}"] = r.get("type", "Роза")
        ws[f"D{rn}"] = r.get("variety", "Mix")
        for lab, price in r.get("prices", {}).items():
            ws[f"{BIF_LENGTH_COL[lab]}{rn}"] = price
        ws[f"M{rn}"] = 1
        ws[f"N{rn}"] = "HB"
        ws[f"O{rn}"] = r["stems"]
        ws[f"P{rn}"] = r.get("sum", "")
    wb.save(path)
    wb.close()
    return path


def make_template(path: Path, lines: list[dict], lengths=("50", "60", "70", "80")) -> Path:
    """lines: {code, title, qtys {length: qty}}; «Итог» считается автоматически."""
    wb = Workbook()
    ws = wb.active
    ws["D2"] = "Пишем в ячейках данные при приеме товара…"
    h = TPL_HEADER_ROW
    ws[f"C{h}"] = "Эквадор"
    for lab in lengths:
        ws[f"{TPL_LENGTH_COL[lab]}{h}"] = lab

    totals = {lab: 0 for lab in lengths}
    for i, ln in enumerate(lines):
        rn = h + 1 + i
        ws[f"C{rn}"] = ln["code"]
        ws[f"D{rn}"] = ln.get("title", ln["code"])
        ws[f"E{rn}"] = i + 1
        for lab, qty in ln.get("qtys", {}).items():
            if qty:
                ws[f"{TPL_LENGTH_COL[lab]}{rn}"] = qty
                totals[lab] += qty
    trn = h + 1 + len(lines)
    ws[f"D{trn}"] = "Итог"
    for lab in lengths:
        ws[f"{TPL_LENGTH_COL[lab]}{trn}"] = totals[lab]
    wb.save(path)
    wb.close()
    return path


def biflorica_rows(path: Path) -> dict[int, dict[str, str]]:
    rows = grid_by_row(read_excel_grid(path))
    return {rn: c for rn, c in rows.items() if rn > BIF_HEADER_ROW}


def mix_row(path: Path, plantation: str) -> dict[str, str] | None:
    for c in biflorica_rows(path).values():
        if c.get("B") == plantation:
            return c
    return None


# === Этап 1: parse_sklad_template ===================================================


@test("1.1 parse_sklad_template: базовый разбор кодов/количеств/итогов")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "Mix R", "title": "Роза Красная Микс", "qtys": {"50": 25, "60": 25}},
        {"code": "Mondial", "title": "Роза Мондиаль", "qtys": {"60": 50}},
    ])
    demand = parse_sklad_template(tpl)
    assert demand.totals == {"50": 25, "60": 75, "70": 0}, demand.totals
    assert len(demand.lines) == 2, demand.lines
    assert demand.lines[0].code == "Mix R"
    assert demand.lines[0].qtys == {"50": 25, "60": 25, "70": 0}
    # "50"/"70" присутствуют в шапке (общая для всех строк) — попадают в qtys нулём,
    # раз в этой конкретной строке значения нет.
    assert demand.lines[1].qtys == {"50": 0, "60": 50, "70": 0}


@test("1.2 parse_sklad_template: строки с нулевым спросом всё равно попадают в lines")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "Candlelight", "title": "Роза Канделайт", "qtys": {}},
        {"code": "Mix 2", "title": "Роза из Документа", "qtys": {"60": 125}},
    ])
    demand = parse_sklad_template(tpl)
    assert len(demand.lines) == 2, demand.lines
    assert demand.lines[0].qtys.get("60", 0) == 0
    assert demand.totals["60"] == 125


@test("1.3 parse_sklad_template: «Итог» останавливает разбор, дальше не читаем")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"50": 10}},
    ])
    # допишем «мусорную» строку после Итога — не должна попасть в lines
    wb_path = tpl
    from openpyxl import load_workbook
    wb = load_workbook(wb_path)
    ws = wb.active
    ws["C12"] = "Junk"
    ws["D12"] = "Мусорная строка после Итога"
    ws["F12"] = 999
    wb.save(wb_path)
    wb.close()
    demand = parse_sklad_template(tpl)
    assert len(demand.lines) == 1, demand.lines
    assert demand.totals["50"] == 10, demand.totals


@test("1.4 parse_sklad_template: нет таблицы «Эквадор» → понятная ошибка")
def _(tmp: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Пустой файл без нужной таблицы"
    p = tmp / "empty.xlsx"
    wb.save(p)
    wb.close()
    try:
        parse_sklad_template(p)
        raise AssertionError("должно было упасть — нет таблицы Эквадор")
    except RuntimeError as e:
        assert "Эквадор" in str(e), e


@test("1.5 parse_sklad_template: сумма строк может не совпадать с «Итог» — берём Итог")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"60": 30}},
    ])
    # руками портим Итог, чтобы не совпадал с суммой строк (30)
    from openpyxl import load_workbook
    wb = load_workbook(tpl)
    ws = wb.active
    ws["G10"] = 999  # G — колонка 60см, строка Итога = header+1+1=10
    wb.save(tpl)
    wb.close()
    demand = parse_sklad_template(tpl)
    assert demand.totals["60"] == 999, demand.totals  # именно Итог, не сумма строк


# === Этап 2: чтение остатка Mix (scan_mix_pool / стебли по длинам) ===================


@test("2.1 scan_mix_pool: обычная (не сплит) строка")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 300},
        {"plantation": "P2", "variety": "Explorer", "prices": {"60": 0.3}, "stems": 500},  # не Mix
    ])
    rows = biflorica_rows(bif)
    header_row, length_map = _biflorica_header(grid_by_row(read_excel_grid(bif)), bif)
    pool = scan_mix_pool(grid_by_row(read_excel_grid(bif)), header_row, length_map, "60")
    assert len(pool) == 1, pool  # не-Mix строка не должна попасть в пул
    assert pool[0].plantation == "P1"
    assert pool[0].stems == 300
    assert abs(pool[0].price - 0.2) < 1e-9


@test("2.2 _stems_for_length: сплит-строка «a|b» — каждая длина получает свою часть")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2, "70": 0.2}, "stems": "300|150"},
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    row = grid[header_row + 1]
    assert _stems_for_length(row, "60", length_map) == 300
    assert _stems_for_length(row, "70", length_map) == 150
    assert _stems_for_length(row, "50", length_map) == 0  # не заявленная длина


@test("2.3 _stems_for_length: тройной сплит «a|b|c»")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix",
         "prices": {"50": 0.5, "60": 0.5, "70": 0.5}, "stems": "40|60|20"},
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    row = grid[header_row + 1]
    assert _stems_for_length(row, "50", length_map) == 40
    assert _stems_for_length(row, "60", length_map) == 60
    assert _stems_for_length(row, "70", length_map) == 20


@test("2.4 _priced_length_labels: порядок по колонкам, а не по цене")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        # 70 (колонка H) стоит раньше 60 (колонка G) в словаре prices, но метки должны
        # выйти в порядке КОЛОНОК: 60, 70
        {"plantation": "P1", "variety": "Mix", "prices": {"70": 0.5, "60": 0.5}, "stems": "1|1"},
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    row = grid[header_row + 1]
    assert _priced_length_labels(row, length_map) == ["60", "70"]


@test("2.5 scan_mix_pool: строки с нулевой/пустой ценой не попадают в пул")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {}, "stems": 100},  # без цены вообще
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    pool = scan_mix_pool(grid, header_row, length_map, "60")
    assert pool == [], pool


@test("2.6 _is_mix_rose: MixAlstromeria и подобные НЕ считаются миксом розы")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "type": "Альстромерия", "variety": "MixAlstromeria",
         "prices": {"60": 0.2}, "stems": 100},
        {"plantation": "P2", "type": "Роза", "variety": "Mix",
         "prices": {"60": 0.2}, "stems": 100},
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    pool = scan_mix_pool(grid, header_row, length_map, "60")
    assert len(pool) == 1, pool
    assert pool[0].plantation == "P2"


# === Этап 3: plan_mix_allocation (распределение) =====================================


@test("3.1 plan_mix_allocation: жадно берёт дорогие строки первыми")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "cheap", "variety": "Mix", "prices": {"60": 0.10}, "stems": 1000},
        {"plantation": "expensive", "variety": "Mix", "prices": {"60": 0.30}, "stems": 50},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    plan = plans[0]
    assert plan.takes[0].plantation == "expensive" and plan.takes[0].take_qty == 50
    assert plan.takes[1].plantation == "cheap" and plan.takes[1].take_qty == 50
    expected_avg = (50 * 0.30 + 50 * 0.10) / 100
    assert abs(plan.avg_price - expected_avg) < 1e-9


@test("3.2 plan_mix_allocation: при равной цене порядок стабилен (по номеру строки)")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "first", "variety": "Mix", "prices": {"60": 0.2}, "stems": 40},
        {"plantation": "second", "variety": "Mix", "prices": {"60": 0.2}, "stems": 40},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 50}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    takes = plans[0].takes
    assert takes[0].plantation == "first" and takes[0].take_qty == 40
    assert takes[1].plantation == "second" and takes[1].take_qty == 10


@test("3.3 plan_mix_allocation: нехватка остатка → RuntimeError с понятным текстом")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 500}}])
    demand = parse_sklad_template(tpl)
    try:
        plan_mix_allocation(demand, bif, log=lambda _m: None)
        raise AssertionError("должно было упасть — не хватает остатка")
    except RuntimeError as e:
        assert "не хватает" in str(e) and "400" in str(e), e


@test("3.4 plan_mix_allocation: нулевой спрос по длине — план не создаётся")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2, "70": 0.2}, "stems": "100|50"},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    assert len(plans) == 1, plans
    assert plans[0].length == "60"


@test("3.5 plan_mix_allocation: точное совпадение (взяли ровно весь остаток)")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    assert plans[0].takes[0].take_qty == 100


# === Этап 4: apply_mix_plans_to_biflorica (запись изменений) ========================


@test("4.1 apply: обычная строка — частичное списание, остаток и сумма верны")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 300},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "P1")
    assert row is not None, "строка P1 не должна исчезнуть"
    assert row["O"] == "200", row
    assert row["P"] == "40", row  # 200 * 0.2


@test("4.2 apply: обычная строка — полное списание удаляет строку")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    assert mix_row(out, "P1") is None, "строка должна быть удалена полностью"


@test("4.3 apply: сплит-строка — списали одну длину, вторая (не в спросе) не тронута")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.25, "70": 0.25}, "stems": "200|150"},
        {"plantation": "P2", "variety": "Mix", "prices": {"60": 0.20}, "stems": 300},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 400}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "P1")
    assert row is not None, "70см в строке не тронуты спросом — строка должна остаться"
    assert "F" not in row and row.get("G") is None or row.get("G") is None, row  # цена 60 очищена
    assert row.get("H") == "0.25", row  # цена 70 осталась
    assert row["O"] == "150", row  # 70см нетронуты
    row2 = mix_row(out, "P2")
    assert row2["O"] == "100", row2  # 300 - 200(взято) = 100


@test("4.4 apply: тройной сплит — одна длина полностью, вторая частично, третья цела")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix",
         "prices": {"50": 0.5, "60": 0.5, "70": 0.5}, "stems": "40|60|20"},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"50": 40}},
        {"code": "B", "title": "B", "qtys": {"60": 30}},
    ])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "P1")
    assert row is not None
    assert row.get("F") is None, row  # 50 полностью взяли — цена очищена
    assert row.get("G") == "0.5", row  # 60 осталась (частично)
    assert row.get("H") == "0.5", row  # 70 не тронута
    assert row["O"] == "30|20", row  # 60-30=30 (остаток), 70=20 (как было)
    assert row["P"] == "25", row  # 30*0.5 + 20*0.5


@test("4.5 apply: обе активные длины (50 и 60) в одной сплит-строке — раздельное списание")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "PX", "variety": "Mix", "prices": {"50": 0.3, "60": 0.3}, "stems": "80|120"},
        {"plantation": "filler50", "variety": "Mix", "prices": {"50": 0.1}, "stems": 1000},
        {"plantation": "filler60", "variety": "Mix", "prices": {"60": 0.1}, "stems": 1000},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A50", "title": "A50", "qtys": {"50": 50}},   # частично из PX (80 → 30 остаток)
        {"code": "A60", "title": "A60", "qtys": {"60": 120}},  # ровно всё из PX
    ])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "PX")
    assert row is not None, "50см ещё осталось — строка должна выжить"
    assert row.get("F") == "0.3", row  # 50 осталась (частично)
    assert row.get("G") is None, row  # 60 полностью взяли — цена очищена
    assert row["O"] == "30", row  # 80 - 50 = 30, единственная оставшаяся длина → просто число


@test("4.6 apply: несколько сплит-строк — не хватает одной, добираем из другой")
def _(tmp: Path) -> None:
    # Ровно сценарий из жалобы: 1025 суммарно на 60см, часть строк — сплит.
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.18}, "stems": 300},
        {"plantation": "P2", "variety": "Mix", "prices": {"60": 0.20, "70": 0.20}, "stems": "300|100"},
        {"plantation": "P3", "variety": "Mix", "prices": {"60": 0.25, "70": 0.25}, "stems": "275|50"},
        {"plantation": "P4", "variety": "Mix", "prices": {"60": 0.19}, "stems": 150},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 500}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)

    rows_after = biflorica_rows(out)
    assert len(rows_after) == 5, rows_after  # 4 исходных Mix + 1 новая позиция

    def stems60(plant: str) -> int:
        row = mix_row(out, plant)
        if row is None or row.get("G") is None:
            return 0
        return int(float(str(row["O"]).split("|")[0]))

    def stems70(plant: str) -> int:
        row = mix_row(out, plant)
        if row is None or row.get("H") is None:
            return 0
        parts = str(row["O"]).split("|")
        return int(float(parts[-1]))

    total_60_left = stems60("P1") + stems60("P2") + stems60("P3") + stems60("P4")
    total_70_left = stems70("P2") + stems70("P3")
    assert total_60_left == 525, total_60_left  # 1025 - 500
    assert total_70_left == 150, total_70_left  # 70см спрос был 0 — не должно пострадать


@test("4.7 apply: у новой позиции цена и сумма согласованы с планом")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.20}, "stems": 100},
        {"plantation": "P2", "variety": "Mix", "prices": {"60": 0.30}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "CodeX", "title": "CodeX", "qtys": {"60": 150}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    new_row = None
    for c in biflorica_rows(out).values():
        if c.get("D") == "CodeX":
            new_row = c
            break
    assert new_row is not None, "новая позиция CodeX не найдена"
    assert new_row["O"] == "150", new_row
    # avg = (100*0.30 + 50*0.20) / 150 = 0.2667
    expected_sum = round(150 * ((100 * 0.30 + 50 * 0.20) / 150), 2)
    assert float(new_row["P"]) == expected_sum, (new_row["P"], expected_sum)


# === Этап 5: путь целиком (run_mix_separation) =======================================


@test("5.1 run_mix_separation: реальный шаблон + архивный Biflorica (регрессия)")
def _(tmp: Path) -> None:
    template = ROOT / "шаблоны тест" / "шаюлон 27,7,26.xlsx"
    bif_src = (
        ROOT / "data/downloads/biflorica/архив/2026-06-03_121632/10738107__2026-05-30.xlsx"
    )
    if not template.is_file() or not bif_src.is_file():
        print("    (пропуск — нет архивных тестовых файлов на этой машине)")
        return
    bif_copy = tmp / "bif_regression.xlsx"
    shutil.copy2(bif_src, bif_copy)
    demand, plans, out = run_mix_separation(
        template_path=template, biflorica_path=bif_copy, log=lambda _m: None
    )
    assert demand.totals["50"] == 100 and demand.totals["60"] == 550, demand.totals
    assert len(plans) == 2, plans


@test("5.2 run_mix_separation: пустой (не заполненный) шаблон без 50/60 → ошибка")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {}}])
    try:
        run_mix_separation(template_path=tpl, biflorica_path=bif, log=lambda _m: None)
        raise AssertionError("должно было упасть — в шаблоне нет заполненных 50/60")
    except RuntimeError as e:
        assert "нет заполненных" in str(e), e


@test("5.3 run_mix_separation: резервная копия «до миксов» реально создаётся")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 50}}])
    run_mix_separation(template_path=tpl, biflorica_path=bif, log=lambda _m: None)
    backups = list(tmp.glob("bif до миксов*.xlsx"))
    assert len(backups) == 1, list(tmp.iterdir())


# === Этап 6: корнер-кейсы ============================================================


@test("6.1 parse_sklad_template: реальный шаблон — берёт ПЕРВУЮ таблицу «Эквадор», не вторую")
def _(tmp: Path) -> None:
    # В реальном шаблоне ДВЕ таблицы с меткой «Эквадор»: роза/микс (строка 8) и
    # гвоздика (строка 19) — обе с теми же колонками длин. Разбор должен остановиться
    # на первой и не приплюсовать гвоздичные позиции к спросу на миксы.
    template = ROOT / "шаблоны тест" / "Шаблон 11.08.26.xlsx"
    if not template.is_file():
        print("    (пропуск — нет архивного шаблона на этой машине)")
        return
    demand = parse_sklad_template(template)
    codes = {ln.code for ln in demand.lines}
    assert "Mix R" in codes and "Mondial" in codes, codes
    assert not any("Гвоздика" in ln.title_ru for ln in demand.lines), [
        ln.title_ru for ln in demand.lines
    ]


@test("6.2 parse_sklad_template: регистр и пробелы в метке «Эквадор» не важны")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 10}}])
    from openpyxl import load_workbook
    wb = load_workbook(tpl)
    ws = wb.active
    ws["C8"] = "  ЭКВАДОР  "  # верхний регистр + пробелы вместо "Эквадор"
    wb.save(tpl)
    wb.close()
    demand = parse_sklad_template(tpl)
    assert demand.totals["60"] == 10, demand.totals


@test("6.3 parse_sklad_template: код и в B, и в C одновременно — код берётся из B")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "FromC", "title": "T", "qtys": {"60": 5}}])
    from openpyxl import load_workbook
    wb = load_workbook(tpl)
    ws = wb.active
    ws["B8"] = "Эквадор"  # метка теперь и в B — приоритет колонки B над C
    ws["B9"] = "FromB"
    wb.save(tpl)
    wb.close()
    demand = parse_sklad_template(tpl)
    assert demand.lines[0].code == "FromB", demand.lines[0].code


@test("6.4 _as_qty: отрицательные и текстовые значения → 0, не падают и не уходят в минус")
def _(tmp: Path) -> None:
    assert _as_qty(-5) == 0
    assert _as_qty("-5") == 0
    assert _as_qty("н/д") == 0
    assert _as_qty("") == 0
    assert _as_qty(None) == 0
    assert _as_qty("12.6") == 13  # обычное округление продолжает работать


@test("6.5 _stems_for_length: пробелы вокруг «|» в сплит-строке не портят разбор")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2, "70": 0.2},
         "stems": " 300 | 150 "},
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    row = grid[header_row + 1]
    assert _stems_for_length(row, "60", length_map) == 300
    assert _stems_for_length(row, "70", length_map) == 150


@test("6.6a apply: сплит-строка 50|60 — ОБЕ целевые длины взяты полностью → строка удаляется")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"50": 0.2, "60": 0.2}, "stems": "100|50"},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"50": 100}},
        {"code": "B", "title": "B", "qtys": {"60": 50}},
    ])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    assert mix_row(out, "P1") is None, "обе целевые длины выбраны без остатка — строка должна исчезнуть"


@test("6.6b apply: сплит-строка 60|80 — 60 выбрано полностью, 80 НЕ целевая длина, не трогаем")
def _(tmp: Path) -> None:
    # 80см никогда не входит в спрос миксов (_TARGET_LENGTHS = 50/60/70) — даже если
    # делить с 60см полностью нечего, 80см в той же строке остаётся как было,
    # строка не удаляется (та же защита из коммита 8122d2a).
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"60": 0.2, "80": 0.2}, "stems": "100|50"},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"60": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "P1")
    assert row is not None, "80см не в спросе — строка не должна исчезать"
    assert row.get("G") is None, row  # цена 60 очищена — эта длина выбрана полностью
    assert row.get("I") == "0.2", row  # цена 80 нетронута
    assert row["O"] == "50", row  # только 80см — единственная оставшаяся длина


@test("6.7 parse_sklad_template: одинаковый код в двух строках — считаются раздельно")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "Dup", "title": "Первая", "qtys": {"60": 20}},
        {"code": "Dup", "title": "Вторая", "qtys": {"60": 30}},
    ])
    demand = parse_sklad_template(tpl)
    assert len(demand.lines) == 2, demand.lines
    assert demand.totals["60"] == 50, demand.totals
    codes = [ln.code for ln in demand.lines]
    assert codes == ["Dup", "Dup"], codes


@test("6.8 parse_sklad_template: таблица «Эквадор» есть, но без 50/60/70 (только 80) → ошибка")
def _(tmp: Path) -> None:
    # Разбор миксов работает только с длинами 50/60/70 — таблица без них не должна
    # молча давать пустой спрос, а должна явно сообщать, что подходящей таблицы нет.
    tpl = make_template(
        tmp / "tpl.xlsx",
        [{"code": "A", "title": "A", "qtys": {"80": 10}}],
        lengths=("80",),
    )
    try:
        parse_sklad_template(tpl)
        raise AssertionError("должно было упасть — нет колонок 50/60")
    except RuntimeError as e:
        assert "50" in str(e) or "60" in str(e), e


# === Этап 7: 70см — полноценная целевая длина (наравне с 50/60) ====================


@test("7.1 parse_sklad_template: спрос по 70 см учитывается в totals и qtys")
def _(tmp: Path) -> None:
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"70": 40}},
    ])
    demand = parse_sklad_template(tpl)
    assert demand.totals["70"] == 40, demand.totals
    assert demand.lines[0].qtys["70"] == 40


@test("7.2 scan_mix_pool: остаток Mix 70 см читается так же, как 50/60")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"70": 0.22}, "stems": 300},
        {"plantation": "P2", "variety": "Explorer", "prices": {"70": 0.30}, "stems": 500},  # не Mix
    ])
    grid = grid_by_row(read_excel_grid(bif))
    header_row, length_map = _biflorica_header(grid, bif)
    pool = scan_mix_pool(grid, header_row, length_map, "70")
    assert len(pool) == 1, pool
    assert pool[0].plantation == "P1" and pool[0].stems == 300


@test("7.3 plan_mix_allocation: 70 см распределяется жадно по цене, как 50/60")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "cheap", "variety": "Mix", "prices": {"70": 0.10}, "stems": 1000},
        {"plantation": "expensive", "variety": "Mix", "prices": {"70": 0.30}, "stems": 50},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"70": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    assert len(plans) == 1 and plans[0].length == "70"
    assert plans[0].takes[0].plantation == "expensive" and plans[0].takes[0].take_qty == 50
    assert plans[0].takes[1].plantation == "cheap" and plans[0].takes[1].take_qty == 50


@test("7.4 plan_mix_allocation: нехватка остатка на 70 см → та же понятная ошибка")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"70": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"70": 500}}])
    demand = parse_sklad_template(tpl)
    try:
        plan_mix_allocation(demand, bif, log=lambda _m: None)
        raise AssertionError("должно было упасть — не хватает остатка на 70см")
    except RuntimeError as e:
        assert "не хватает" in str(e) and "400" in str(e), e


@test("7.5 apply: полное списание строки 70 см удаляет её, как для 50/60")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix", "prices": {"70": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [{"code": "A", "title": "A", "qtys": {"70": 100}}])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    assert mix_row(out, "P1") is None


@test("7.6 apply: все три длины 50/60/70 в одной сплит-строке разбираются по отдельности")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P1", "variety": "Mix",
         "prices": {"50": 0.4, "60": 0.4, "70": 0.4}, "stems": "40|60|20"},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"50": 40}},   # выбрана полностью
        {"code": "B", "title": "B", "qtys": {"60": 30}},   # частично (было 60)
        {"code": "C", "title": "C", "qtys": {"70": 20}},   # выбрана полностью
    ])
    demand = parse_sklad_template(tpl)
    plans = plan_mix_allocation(demand, bif, log=lambda _m: None)
    assert {p.length for p in plans} == {"50", "60", "70"}, [p.length for p in plans]
    out = apply_mix_plans_to_biflorica(bif, plans, log=lambda _m: None)
    row = mix_row(out, "P1")
    assert row is not None, "60см осталось 30 — строка должна выжить"
    assert row.get("F") is None, row  # 50 выбрана полностью — цена очищена
    assert row.get("G") == "0.4", row  # 60 частично осталась
    assert row.get("H") is None, row  # 70 выбрана полностью — цена очищена
    assert row["O"] == "30", row  # единственная оставшаяся длина — просто число


@test("7.7 run_mix_separation: реальный вертикальный срез — 70см вместе с 50/60 в одном прогоне")
def _(tmp: Path) -> None:
    bif = make_biflorica(tmp / "bif.xlsx", [
        {"plantation": "P50", "variety": "Mix", "prices": {"50": 0.2}, "stems": 100},
        {"plantation": "P60", "variety": "Mix", "prices": {"60": 0.2}, "stems": 100},
        {"plantation": "P70", "variety": "Mix", "prices": {"70": 0.2}, "stems": 100},
    ])
    tpl = make_template(tmp / "tpl.xlsx", [
        {"code": "A", "title": "A", "qtys": {"50": 100}},
        {"code": "B", "title": "B", "qtys": {"60": 100}},
        {"code": "C", "title": "C", "qtys": {"70": 100}},
    ])
    demand, plans, out = run_mix_separation(template_path=tpl, biflorica_path=bif, log=lambda _m: None)
    assert {p.length for p in plans} == {"50", "60", "70"}, [p.length for p in plans]
    for plant in ("P50", "P60", "P70"):
        assert mix_row(out, plant) is None, f"{plant} должен был полностью разобраться"
    new_codes = {c.get("D") for c in biflorica_rows(out).values() if c.get("D") in {"A", "B", "C"}}
    assert new_codes == {"A", "B", "C"}, new_codes


# === Раннер ===========================================================================


def main() -> int:
    base = Path(tempfile.mkdtemp(prefix="cvetopt-mix-tests-"))
    passed = 0
    failed: list[str] = []
    for i, (name, fn) in enumerate(TESTS):
        tmp = base / f"case_{i:02d}"
        tmp.mkdir(parents=True, exist_ok=True)
        try:
            fn(tmp)
        except Exception as e:  # noqa: BLE001
            failed.append(name)
            print(f"FAIL  {name}")
            print(f"      {type(e).__name__}: {e}")
            traceback.print_exc(limit=3)
            print()
        else:
            passed += 1
            print(f"PASS  {name}")
    shutil.rmtree(base, ignore_errors=True)

    print()
    print(f"Итого: {passed}/{len(TESTS)} прошли")
    if failed:
        print("Упавшие:")
        for name in failed:
            print(f"  - {name}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
