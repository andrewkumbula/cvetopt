#!/usr/bin/env python3
"""
Сквозная проверка группового скачивания Biflorica (несколько заказов с одной датой
вылета → один файл) — не unit-тест отдельных функций, а настоящий Playwright-браузер
против локального фейкового сервера с реальной вёрсткой портала (id="list-orders-*",
класс кнопки group-orders-btn и т.д. — сверено с "пример страницы заказы/Заказы -
Biflorica.com.html").

Живой Biflorica — Angular-приложение: при открытии страницы оно само перезапрашивает
список заказов с /api/orders и без реального бэкенда стирает любую подставленную
разметку. Полностью повторить это без сервера Biflorica нельзя. Но сам код скачивания
(_collect_all_orders_paginated, _download_group_report, _walk_pages_setting_checkboxes
из cvetopt.scrapers.biflorica) ничего не знает про Angular — он работает с обычным DOM
через те же CSS-селекторы, что и в проде. Этот фейковый сервер отдаёт ту же разметку
через простой JS (без Angular) и реально стримит xlsx по клику — этого достаточно,
чтобы прогнать настоящий код скачивания и проверить, что чекбоксы двух заказов с
одной датой вылета реально уходят ОДНИМ запросом и сохраняются в ОДИН файл.

Не покрыто (нужен реальный портал): поведение самого Angular-приложения Biflorica
(его JS-логика показа/скрытия кнопки, реальная генерация xlsx их сервером).

Запуск: .venv/bin/python scripts/biflorica_group_download_e2e_local.py
"""

from __future__ import annotations

import asyncio
import shutil
import sys
import tempfile
import threading
from datetime import date
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import uvicorn  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from fastapi.responses import HTMLResponse, StreamingResponse  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from playwright.async_api import async_playwright  # noqa: E402

import cvetopt.scrapers.biflorica as bif  # noqa: E402
from cvetopt.core.settings import AppYamlConfig, BifloricaPortalConfig, SelectionOverride  # noqa: E402
from cvetopt.invoice.biflorica_split import diagnose_biflorica_report  # noqa: E402
from cvetopt.invoice.xlsx_read import grid_by_row, read_excel_grid  # noqa: E402

PORT = 8936
BASE = f"http://127.0.0.1:{PORT}"

# Заказы: два с одной датой вылета (воспроизводит жалобу — «12 сентября было два
# файла»), один с другой — контроль, что он не подмешивается в группу.
FAKE_ORDERS: dict[str, dict] = {
    "10724085": {"created": "09 May 2026, 04:02", "flight": "09 May 2026"},
    "10724268": {"created": "09 May 2026, 10:32", "flight": "09 May 2026"},
    "10720147": {"created": "30 Apr 2026, 09:02", "flight": "02 May 2026"},
}
DOWNLOAD_LOG: list[list[str]] = []

fake_app = FastAPI()


@fake_app.get("/marketPlaceNew/default/orders", response_class=HTMLResponse)
def _orders_page() -> str:
    rows = []
    for oid, o in FAKE_ORDERS.items():
        rows.append(f"""
        <tr class="buyer">
          <td>
            <input type="checkbox" class="checkbox order-checkbox" id="list-orders-{oid}">
            <label for="list-orders-{oid}">{oid}</label>
          </td>
          <td>{o['created']}</td>
          <td>SO_Ecuador</td>
          <td class="ng-hide"></td>
          <td>{o['flight']}</td>
          <td>D&amp;C (BiFlorica.Ecuador) — VOSTOK</td>
        </tr>""")
    rows_html = "\n".join(rows)
    # Разметка сверена с "пример страницы заказы/Заказы - Biflorica.com.html":
    # id="list-orders-<id>", label[for=...], button.bf-btn-copy-ico.group-orders-btn.
    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Заказы (fake)</title></head>
<body>
<div id="orderController">
  <label><input type="radio" id="all-requests" checked> Все</label>
  <table class="bf-table-with-grey-border">
    <tbody>{rows_html}</tbody>
  </table>
  <button class="bf-btn-copy-ico group-orders-btn" title="Отчет по сделкам" disabled
          onclick="doDownload()">Отчет по сделкам</button>
</div>
<script>
function updateButton() {{
  var any = Array.prototype.some.call(
    document.querySelectorAll('.order-checkbox'), function(x) {{ return x.checked; }});
  document.querySelector('.group-orders-btn').disabled = !any;
}}
document.querySelectorAll('.order-checkbox').forEach(function(cb) {{
  cb.addEventListener('change', updateButton);
}});
function doDownload() {{
  var ids = Array.prototype.filter.call(
    document.querySelectorAll('.order-checkbox'), function(x) {{ return x.checked; }}
  ).map(function(x) {{ return x.id.replace('list-orders-', ''); }});
  window.location.href = '/download?ids=' + ids.join(',');
}}
</script>
</body></html>"""


@fake_app.get("/download")
def _download(ids: str) -> StreamingResponse:
    order_ids = ids.split(",")
    DOWNLOAD_LOG.append(order_ids)

    wb = Workbook()
    ws = wb.active
    h = 6
    ws[f"A{h}"] = "ДАТА И ВРЕМЯ СДЕЛКИ"
    ws[f"B{h}"] = "ПЛАНТАЦИЯ"
    ws[f"C{h}"] = "ТИП"
    ws[f"D{h}"] = "СОРТ"
    ws[f"G{h}"] = "60"
    ws[f"O{h}"] = "ВСЕГО СТЕБЛЕЙ"
    ws[f"P{h}"] = "СУММА СДЕЛКИ"
    for i, oid in enumerate(order_ids):
        rn = h + 1 + i
        ws[f"A{rn}"] = "09 May 2026, 10:32"
        ws[f"B{rn}"] = f"PLANT-{oid}"
        ws[f"C{rn}"] = "Роза"
        ws[f"D{rn}"] = "Mix"
        ws[f"G{rn}"] = 0.2
        ws[f"O{rn}"] = 100
        ws[f"P{rn}"] = 20
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"BiFlorica-{'+'.join(order_ids)}__2026-05-09.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _run_fake_server() -> None:
    uvicorn.run(fake_app, host="127.0.0.1", port=PORT, log_level="warning")


async def _amain() -> int:
    threading.Thread(target=_run_fake_server, daemon=True).start()
    await asyncio.sleep(0.8)

    portal = BifloricaPortalConfig(orders_url=f"{BASE}/marketPlaceNew/default/orders")
    yaml_cfg = AppYamlConfig()
    override = SelectionOverride(min_age_days=0, max_age_days=200)
    today = date(2026, 9, 18)
    checks: list[tuple[str, bool, str]] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        checks.append((name, ok, detail))
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f" — {detail}" if detail and not ok else ""))

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page()

        async def lg(msg: str) -> None:
            pass

        await page.goto(portal.orders_url, wait_until="domcontentloaded")
        await bif._await_orders_table(page, portal, lg)
        await bif._ensure_tab_all(page, portal, lg)
        orders = await bif._collect_all_orders_paginated(page, portal, yaml_cfg, today, override, lg)
        check("собраны все 3 заказа", len(orders) == 3, f"нашли {len(orders)}")

        groups: dict[date, list] = {}
        for o in orders:
            groups.setdefault(o.flight_date, []).append(o)
        target_date = date(2026, 5, 9)
        group_ids = [o.order_id for o in groups.get(target_date, [])]
        check(
            "группировка по дате вылета нашла пару 10724085+10724268",
            set(group_ids) == {"10724085", "10724268"},
            str(group_ids),
        )
        other_date_ids = {o.order_id for o in groups.get(date(2026, 5, 2), [])}
        check(
            "заказ с другой датой (10720147) — своя отдельная группа",
            other_date_ids == {"10720147"},
            str(other_date_ids),
        )

        if group_ids:
            tmp = Path(tempfile.mkdtemp())
            dest = tmp / f"BiFlorica-{'+'.join(sorted(group_ids))}__{target_date.isoformat()}.xlsx"
            DOWNLOAD_LOG.clear()
            await bif._download_group_report(page, portal, group_ids, dest, lg)

            check(
                "сервер получил РОВНО ОДИН запрос на скачивание группы",
                len(DOWNLOAD_LOG) == 1,
                str(DOWNLOAD_LOG),
            )
            check(
                "в этом запросе оба id вместе, не по одному",
                bool(DOWNLOAD_LOG) and set(DOWNLOAD_LOG[0]) == set(group_ids),
                str(DOWNLOAD_LOG[0]) if DOWNLOAD_LOG else "нет запросов",
            )
            check("файл сохранён на диск", dest.exists())
            if dest.exists():
                bad = diagnose_biflorica_report(dest)
                check("файл проходит diagnose_biflorica_report", bad is None, str(bad))
                rows = grid_by_row(read_excel_grid(dest))
                plantations = {c.get("B") for c in rows.values() if str(c.get("B", "")).startswith("PLANT-")}
                expected = {f"PLANT-{oid}" for oid in group_ids}
                check("в объединённом файле обе сделки, не одна", plantations == expected, str(plantations))
                shutil.rmtree(tmp, ignore_errors=True)

            for oid in group_ids:
                checked = await page.locator(f"#list-orders-{oid}").is_checked()
                check(f"галочка {oid} снята после скачивания", not checked)

        await browser.close()

    passed = sum(1 for _, ok, _ in checks if ok)
    print(f"\nИтого: {passed}/{len(checks)} прошли")
    return 0 if passed == len(checks) else 1


def main() -> int:
    return asyncio.run(_amain())


if __name__ == "__main__":
    raise SystemExit(main())
